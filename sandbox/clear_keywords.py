"""Extract keywords from DOCX files."""

import re
import sys

from docx import Document


def extract_words_from_docx(input_path, output_path):
    """Extract words from DOCX and generate keyword phrases."""
    doc = Document(input_path)
    words = set()

    word_pattern = re.compile(r"[a-zA-Zа-яА-ЯәғқңөұүһіӘҒҚҢӨҰҮҺІёЁ]+", re.UNICODE)

    # Извлечение из параграфов
    for para in doc.paragraphs:
        found = word_pattern.findall(para.text)
        words.update(found)

    # Извлечение из таблиц
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                found = word_pattern.findall(cell.text)
                words.update(found)

    # Генерация вариантов с добавлением локаций
    result_lines = []
    for word in sorted(words, key=str.lower):
        result_lines.append(f"{word} казахстан")
        result_lines.append(f"{word} алматы казахстан")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(result_lines))

    print(
        f"Извлечено {len(words)} слов, "
        f"создано {len(result_lines)} ключевых фраз в {output_path}"
    )


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Использование: python script.py input.docx [output.txt]")
        sys.exit(1)

    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else "keywords.txt"

    extract_words_from_docx(input_file, output_file)
